"""Monta o .pptx a partir de um arquivo de conteúdo em Markdown.

Uso:
    python slides/scripts/build_deck.py slides/conteudo/aula-1.md

Gera slides/saida/<nome>.pptx e regenera as figuras citadas, se faltarem.

Gramática do arquivo de conteúdo
--------------------------------
Slides são separados por uma linha com `---`. Dentro de cada slide:

    # Título                  título de slide de conteúdo
    ## Título                 divisória de bloco (slide em negativo)
    ### Subtítulo             linha de apoio sob o título
    @capa                     slide de abertura (usa os dados de estilo.py)
    @encerramento             slide final
    !! Frase                  âncora prática, em destaque
    :fig nome-da-figura       figura gerada por figuras.py
    :fonte Texto              crédito da figura ou da tabela
    :bloco Rótulo             rótulo de bloco no rodapé, vale até mudar
    - item                    marcador de primeiro nível
      - item                  marcador de segundo nível
    > nota                    nota do apresentador (não aparece projetada)

Linhas em branco são ignoradas. Nada de sintaxe além dessa.
"""

from __future__ import annotations

import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(Path(__file__).parent))
import estilo as E  # noqa: E402
import figuras  # noqa: E402

RAIZ = Path(__file__).resolve().parents[1]
DIR_FIGURAS = RAIZ / "build" / "figuras"
DIR_SAIDA = RAIZ / "saida"


def rgb(hexa: str) -> RGBColor:
    return RGBColor.from_string(hexa.lstrip("#").upper())


@dataclass
class Slide:
    tipo: str = "conteudo"  # conteudo | secao | capa | encerramento | ancora
    titulo: str = ""
    subtitulo: str = ""
    bullets: list[tuple[int, str]] = field(default_factory=list)
    figura: str = ""
    fonte: str = ""
    ancora: str = ""
    bloco: str = ""
    notas: list[str] = field(default_factory=list)


def parse(texto: str) -> list[Slide]:
    slides: list[Slide] = []
    bloco_atual = ""
    for bruto in texto.split("\n---\n"):
        linhas = [linha.rstrip() for linha in bruto.strip("\n").split("\n")]
        if not any(linha.strip() for linha in linhas):
            continue
        s = Slide(bloco=bloco_atual)
        for linha in linhas:
            crua, limpa = linha, linha.strip()
            if not limpa:
                continue
            if limpa == "@capa":
                s.tipo = "capa"
            elif limpa == "@encerramento":
                s.tipo = "encerramento"
            elif limpa.startswith("## "):
                s.tipo, s.titulo = "secao", limpa[3:].strip()
            elif limpa.startswith("### "):
                s.subtitulo = limpa[4:].strip()
            elif limpa.startswith("# "):
                s.titulo = limpa[2:].strip()
            elif limpa.startswith("!! "):
                s.tipo, s.ancora = "ancora", limpa[3:].strip()
            elif limpa.startswith(":fig "):
                s.figura = limpa[5:].strip()
            elif limpa.startswith(":fonte "):
                s.fonte = limpa[7:].strip()
            elif limpa.startswith(":bloco "):
                bloco_atual = s.bloco = limpa[7:].strip()
            elif limpa.startswith("> "):
                s.notas.append(limpa[2:].strip())
            elif limpa.startswith("- "):
                nivel = 1 if re.match(r"^\s{2,}- ", crua) else 0
                s.bullets.append((nivel, limpa[2:].strip()))
            else:
                raise SystemExit(f"linha não reconhecida no conteúdo:\n  {linha}")
        slides.append(s)
    return slides


# --- Primitivas de desenho ---------------------------------------------------


def fundo(slide, cor: str):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(cor)


def caixa_texto(slide, x, y, w, h, alinhamento=PP_ALIGN.LEFT, ancora=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ancora
    tf.paragraphs[0].alignment = alinhamento
    return tf


def escreve(par, texto, *, fonte, tamanho, cor, negrito=False, italico=False, espaco_antes=0, espaco_depois=0, entrelinha=1.0):
    """Escreve um parágrafo.

    Trechos entre `**` viram negrito; trechos entre crases viram monoespaçado,
    ligeiramente menor para compensar a largura da fonte de código.
    """
    par.space_before = Pt(espaco_antes)
    par.space_after = Pt(espaco_depois)
    par.line_spacing = entrelinha
    ultimo = None
    for parte in re.split(r"(\*\*.+?\*\*|`[^`]+`)", texto):
        if not parte:
            continue
        em_negrito, mono = negrito, False
        if parte.startswith("**") and parte.endswith("**"):
            parte, em_negrito = parte[2:-2], True
        elif parte.startswith("`") and parte.endswith("`"):
            parte, mono = parte[1:-1], True
        run = par.add_run()
        run.text = parte
        run.font.name = E.FONTE_MONO if mono else fonte
        run.font.size = Pt(tamanho * 0.92 if mono else tamanho)
        run.font.color.rgb = rgb(cor)
        run.font.bold = em_negrito
        run.font.italic = italico
        ultimo = run
    return ultimo


def linha_horizontal(slide, x, y, w, cor, espessura=1.25):
    from pptx.enum.shapes import MSO_SHAPE

    forma = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Pt(espessura))
    forma.fill.solid()
    forma.fill.fore_color.rgb = rgb(cor)
    forma.line.fill.background()
    forma.shadow.inherit = False
    return forma


def rodape(slide, s: Slide, numero: int):
    if s.tipo in ("capa", "secao"):
        return
    tf = caixa_texto(slide, E.MARGEM, E.RODAPE_Y, E.LARGURA - 2 * E.MARGEM - 0.6, 0.34)
    esquerda = " · ".join(x for x in (E.AUTOR, s.bloco) if x)
    escreve(tf.paragraphs[0], esquerda, fonte=E.FONTE_CORPO, tamanho=E.PT_RODAPE, cor=E.TEXTO_FRACO)

    tf_num = caixa_texto(slide, E.LARGURA - E.MARGEM - 0.6, E.RODAPE_Y, 0.6, 0.34, PP_ALIGN.RIGHT)
    escreve(tf_num.paragraphs[0], str(numero), fonte=E.FONTE_CORPO, tamanho=E.PT_RODAPE, cor=E.TEXTO_FRACO)


def notas(slide, s: Slide):
    if not s.notas:
        return
    slide.notes_slide.notes_text_frame.text = "\n".join(s.notas)


# --- Layouts -----------------------------------------------------------------


def desenha_capa(slide, s: Slide):
    fundo(slide, E.FUNDO_SECAO)
    linha_horizontal(slide, E.MARGEM, 2.55, 2.2, E.CAMADA_C, 2.5)

    tf = caixa_texto(slide, E.MARGEM, 2.85, E.LARGURA - 2 * E.MARGEM, 2.0)
    escreve(
        tf.paragraphs[0],
        s.titulo,
        fonte=E.FONTE_TITULO,
        tamanho=E.PT_TITULO_CAPA,
        cor=E.TEXTO_NEGATIVO,
        entrelinha=1.15,
    )
    if s.subtitulo:
        p = tf.add_paragraph()
        escreve(
            p,
            s.subtitulo,
            fonte=E.FONTE_CORPO,
            tamanho=E.PT_SUBTITULO_CAPA,
            cor=E.ACENTO_CLARO,
            espaco_antes=14,
        )

    tf2 = caixa_texto(slide, E.MARGEM, 5.35, E.LARGURA - 2 * E.MARGEM, 1.4)
    escreve(tf2.paragraphs[0], E.AUTOR, fonte=E.FONTE_CORPO, tamanho=16, cor=E.TEXTO_NEGATIVO)
    for linha in (E.INSTITUICAO, E.EVENTO):
        p = tf2.add_paragraph()
        escreve(p, linha, fonte=E.FONTE_CORPO, tamanho=12.5, cor=E.ACENTO_CLARO, espaco_antes=3)


def desenha_secao(slide, s: Slide):
    fundo(slide, E.FUNDO_SECAO)
    linha_horizontal(slide, E.MARGEM, 3.05, 1.6, E.CAMADA_C, 2.5)
    tf = caixa_texto(slide, E.MARGEM, 3.35, E.LARGURA - 2 * E.MARGEM, 1.6)
    escreve(tf.paragraphs[0], s.titulo, fonte=E.FONTE_TITULO, tamanho=32, cor=E.TEXTO_NEGATIVO)
    if s.subtitulo:
        p = tf.add_paragraph()
        escreve(p, s.subtitulo, fonte=E.FONTE_CORPO, tamanho=16, cor=E.ACENTO_CLARO, espaco_antes=10)


def desenha_ancora(slide, s: Slide):
    fundo(slide, E.FUNDO_CAIXA)
    linha_horizontal(slide, E.MARGEM, 2.55, 1.6, E.CAMADA_C, 2.5)
    tf = caixa_texto(slide, E.MARGEM, 2.9, E.LARGURA - 2 * E.MARGEM, 2.6, ancora=MSO_ANCHOR.TOP)
    escreve(
        tf.paragraphs[0],
        s.ancora,
        fonte=E.FONTE_TITULO,
        tamanho=E.PT_ANCORA,
        cor=E.TEXTO,
        italico=True,
        entrelinha=1.35,
    )
    if s.titulo:
        p = tf.add_paragraph()
        escreve(p, s.titulo, fonte=E.FONTE_CORPO, tamanho=13, cor=E.TEXTO_FRACO, espaco_antes=18)


def cabecalho(slide, s: Slide):
    tf = caixa_texto(slide, E.MARGEM, E.TOPO_TITULO, E.LARGURA - 2 * E.MARGEM, 0.9)
    escreve(tf.paragraphs[0], s.titulo, fonte=E.FONTE_TITULO, tamanho=E.PT_TITULO, cor=E.TEXTO)
    y = E.TOPO_TITULO + 0.72
    if s.subtitulo:
        tf2 = caixa_texto(slide, E.MARGEM, y, E.LARGURA - 2 * E.MARGEM, 0.44)
        escreve(tf2.paragraphs[0], s.subtitulo, fonte=E.FONTE_CORPO, tamanho=E.PT_SUBTITULO, cor=E.TEXTO_FRACO)
        y += 0.42
    linha_horizontal(slide, E.MARGEM, y + 0.12, E.LARGURA - 2 * E.MARGEM, E.LINHA, 1.0)
    return y + 0.42


def desenha_bullets(slide, s: Slide, x, y, largura):
    if not s.bullets:
        return
    tf = caixa_texto(slide, x, y, largura, E.RODAPE_Y - y - 0.2)
    primeiro = True
    for nivel, texto in s.bullets:
        par = tf.paragraphs[0] if primeiro else tf.add_paragraph()
        primeiro = False
        par.level = nivel
        marcador = "—  " if nivel == 0 else "·  "
        escreve(
            par,
            marcador + texto,
            fonte=E.FONTE_CORPO,
            tamanho=E.PT_BULLET if nivel == 0 else E.PT_BULLET_2,
            cor=E.TEXTO if nivel == 0 else E.TEXTO_FRACO,
            espaco_antes=0 if primeiro else (13 if nivel == 0 else 6),
            entrelinha=1.22,
        )
        par.space_before = Pt(13 if nivel == 0 else 5)
    tf.paragraphs[0].space_before = Pt(0)


def insere_figura(slide, nome, x, y, largura_max, altura_max, fonte_credito=""):
    caminho = DIR_FIGURAS / f"{nome}.png"
    if not caminho.exists():
        figuras.gerar([nome])
    from PIL import Image

    with Image.open(caminho) as img:
        prop = img.width / img.height
    largura = largura_max
    altura = largura / prop
    if altura > altura_max:
        altura = altura_max
        largura = altura * prop
    esquerda = x + (largura_max - largura) / 2
    topo = y + (altura_max - altura) / 2
    slide.shapes.add_picture(str(caminho), Inches(esquerda), Inches(topo), Inches(largura), Inches(altura))
    if fonte_credito:
        # Crédito logo abaixo da imagem, não no fim da área reservada: colado
        # na figura, é lido como legenda; solto, parece rodapé.
        tf = caixa_texto(slide, esquerda, min(topo + altura + 0.06, E.RODAPE_Y - 0.34), largura, 0.3)
        escreve(tf.paragraphs[0], fonte_credito, fonte=E.FONTE_CORPO, tamanho=E.PT_LEGENDA, cor=E.TEXTO_FRACO, italico=True)


def desenha_conteudo(slide, s: Slide):
    fundo(slide, E.FUNDO)
    y = cabecalho(slide, s)
    util = E.LARGURA - 2 * E.MARGEM
    altura_util = E.RODAPE_Y - y - 0.45

    if s.figura and s.bullets:
        largura_texto = util * 0.42
        desenha_bullets(slide, s, E.MARGEM, y, largura_texto - 0.25)
        insere_figura(slide, s.figura, E.MARGEM + largura_texto, y, util - largura_texto, altura_util, s.fonte)
    elif s.figura:
        insere_figura(slide, s.figura, E.MARGEM, y, util, altura_util, s.fonte)
    else:
        desenha_bullets(slide, s, E.MARGEM, y, util)


def desenha_encerramento(slide, s: Slide):
    fundo(slide, E.FUNDO_SECAO)
    linha_horizontal(slide, E.MARGEM, 2.35, 1.6, E.CAMADA_C, 2.5)
    tf = caixa_texto(slide, E.MARGEM, 2.7, E.LARGURA - 2 * E.MARGEM, 3.4)
    escreve(tf.paragraphs[0], s.titulo, fonte=E.FONTE_TITULO, tamanho=30, cor=E.TEXTO_NEGATIVO)
    for nivel, texto in s.bullets:
        p = tf.add_paragraph()
        escreve(
            p,
            texto,
            fonte=E.FONTE_CORPO,
            tamanho=15 if nivel == 0 else 13,
            cor=E.TEXTO_NEGATIVO if nivel == 0 else E.ACENTO_CLARO,
            espaco_antes=12,
            entrelinha=1.2,
        )


# --- Montagem ----------------------------------------------------------------


def monta(slides: list[Slide], destino: Path):
    prs = Presentation()
    prs.slide_width = Emu(int(E.LARGURA * 914400))
    prs.slide_height = Emu(int(E.ALTURA * 914400))
    vazio = prs.slide_layouts[6]

    numero = 0
    for s in slides:
        slide = prs.slides.add_slide(vazio)
        if s.tipo == "capa":
            desenha_capa(slide, s)
        elif s.tipo == "secao":
            desenha_secao(slide, s)
        elif s.tipo == "ancora":
            numero += 1
            desenha_ancora(slide, s)
        elif s.tipo == "encerramento":
            desenha_encerramento(slide, s)
        else:
            numero += 1
            desenha_conteudo(slide, s)
        rodape(slide, s, numero)
        notas(slide, s)

    destino.parent.mkdir(parents=True, exist_ok=True)
    prs.save(destino)
    return len(slides)


def exporta_pdf(pptx: Path) -> Path | None:
    """Converte para PDF com o PowerPoint instalado. O PDF é o artefato de
    conferência: é nele que se revisa o deck página a página, sem depender de
    abrir o PowerPoint."""
    import subprocess

    script = Path(__file__).parent / "exportar_pdf.ps1"
    pdf = RAIZ / "build" / f"{pptx.stem}.pdf"
    resultado = subprocess.run(
        [
            "powershell",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(script),
            "-Pptx",
            str(pptx),
            "-Pdf",
            str(pdf),
        ],
        capture_output=True,
        text=True,
        timeout=600,
    )
    if resultado.returncode != 0 or not pdf.exists():
        print("aviso: não foi possível gerar o PDF de conferência.")
        detalhe = (resultado.stderr or resultado.stdout).strip().splitlines()
        if detalhe:
            print(f"       {detalhe[0]}")
        print("       Feche o PowerPoint, se estiver aberto, e rode de novo.")
        return None
    return pdf


def main(argv):
    if not argv:
        raise SystemExit("uso: python slides/scripts/build_deck.py <arquivo-de-conteudo.md> [--sem-pdf]")
    sem_pdf = "--sem-pdf" in argv
    origem = Path([a for a in argv if not a.startswith("--")][0])
    if not origem.exists():
        raise SystemExit(f"conteúdo não encontrado: {origem}")
    slides = parse(origem.read_text(encoding="utf-8"))
    destino = DIR_SAIDA / f"{origem.stem}.pptx"
    total = monta(slides, destino)
    print(f"{total} slides gerados em {destino}")
    if not sem_pdf:
        pdf = exporta_pdf(destino)
        if pdf:
            print(f"PDF de conferência em {pdf}")


if __name__ == "__main__":
    main(sys.argv[1:])
